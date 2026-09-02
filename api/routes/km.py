"""api/routes/km.py — the aligned Knowledge Base API.

Serves the KM layer (pipeline/kb/km_query) to the Knowledge Base view: contract families,
the aligned per-field CURRENT value (with superseded history + evidence), and deterministic
aggregation. CLEARANCE-GATED (confidential + admin, from the login session) — this is a
raw knowledge surface; default accounts use chat, where per-retrieval filtering applies.
"""
from __future__ import annotations

import re

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse, Response

from pipeline.config import Config
from pipeline.kb import km_query

from .auth import require_clearance

router = APIRouter(prefix="/km", tags=["km"],
                   dependencies=[Depends(require_clearance)])

_CFG = Config.load()


@router.get("/families")
def list_families(user: dict = Depends(require_clearance)) -> dict:
    """Every contract family with its documents + field counts (richest first)."""
    try:
        return {"families": km_query.families(_CFG, user["role"])}
    except Exception as exc:  # noqa: BLE001 — surface a clean 503 if the store is down
        raise HTTPException(503, f"knowledge base unavailable: {exc}") from exc


@router.get("/family/{group}")
def family_view(group: str, user: dict = Depends(require_clearance)) -> dict:
    """A family's aligned knowledge: canonical parties, the document DAG, and every field's
    current value + superseded history, filtered for the session role."""
    view = km_query.aligned_view(_CFG, group, user["role"])
    if not view["categories"] and not view["parties"]:
        raise HTTPException(404, f"no knowledge-base data for family {group!r} — "
                                 "run: python -m scripts.build_km")
    return view


@router.get("/aggregate")
def aggregate(field_key: str, group: str | None = None, op: str = "sum",
              user: dict = Depends(require_clearance)) -> dict:
    """Deterministic sum/list of a field's numeric values across current statements."""
    if op not in ("sum", "list"):
        raise HTTPException(400, "op must be 'sum' or 'list'")
    return km_query.aggregate(_CFG, field_key, group, user["role"], op)


@router.get("/export")
def export_xlsx(group: str | None = None,
                user: dict = Depends(require_clearance)) -> Response:
    """The aligned knowledge as an Excel workbook: one sheet per contract
    family, one row per current value (category, field, value, trust, source
    document and page), with superseded history alongside. RBAC-filtered with
    the same rules as the on-screen view."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    fams = km_query.families(_CFG, user["role"])
    titles = {f["group"]: f["title"] for f in fams}
    groups = [group] if group else [f["group"] for f in fams]

    wb = Workbook()
    wb.remove(wb.active)
    headers = ["Category", "Field", "Current value", "Unit", "Trust",
               "Verified", "Source document", "Page", "Superseded history"]
    widths = [24, 34, 40, 10, 16, 10, 44, 7, 60]
    used_names: set[str] = set()
    for grp in groups:
        try:
            view = km_query.aligned_view(_CFG, grp, user["role"])
        except Exception:  # noqa: BLE001 — skip a broken family, keep the export
            continue
        if not view["categories"]:
            continue
        # Sheet names: Excel forbids []:*?/\ and caps at 31 chars.
        name = re.sub(r"[\[\]:*?/\\]", " ", str(titles.get(grp, grp)))[:31].strip() or grp[:31]
        n, i = name, 2
        while n in used_names:
            n = f"{name[:28]} {i}"
            i += 1
        used_names.add(n)
        ws = wb.create_sheet(n)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.freeze_panes = "A2"
        for cat in view["categories"]:
            for f in cat["fields"]:
                history = " | ".join(
                    f"{s['value']} ({s['doc_title']})" for s in f["superseded"])
                if not f["current"]:
                    ws.append([cat["category"].replace("_", " "), f["title"],
                               "not stated", "", "", "", "", "", history])
                    continue
                for s in f["current"]:
                    ws.append([
                        cat["category"].replace("_", " "), f["title"],
                        str(s["value"]), s.get("unit") or "",
                        "human verified" if s.get("verified") else "AI extracted",
                        "yes" if s.get("verified") else "",
                        s.get("doc_title") or s["doc_id"], s.get("page") or "",
                        history,
                    ])
    if not wb.sheetnames:  # an xlsx must contain at least one sheet
        wb.create_sheet("Knowledge").append(["No knowledge-base data yet."])

    buf = BytesIO()
    wb.save(buf)
    fname = "knowledge_export.xlsx" if not group else \
        f"knowledge_{re.sub(r'[^A-Za-z0-9_-]+', '_', group)[:40]}.xlsx"
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@router.get("/page/{doc_id}/{page_no}")
def get_page(doc_id: str, page_no: int) -> FileResponse:
    """A page image for the Knowledge Base evidence pane. Mirrors /review/page,
    which is admin-only — Knowledge users (clearance) need their own image route."""
    if not doc_id.isalnum():
        raise HTTPException(400, "bad doc_id")
    if page_no < 1 or page_no > 9999:
        raise HTTPException(400, "bad page number")
    path = _CFG.storage_root / "pages" / doc_id / f"p_{page_no:03d}.png"
    if not path.exists():
        raise HTTPException(404, "page not found")
    return FileResponse(path, media_type="image/png",
                        headers={"Cache-Control": "public, max-age=3600"})
